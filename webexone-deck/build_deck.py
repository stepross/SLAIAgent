"""Build the WebexOne 2026 technical session deck on the official template."""
import os
import re

import cairosvg
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TEMPLATE = "/home/ubuntu/.cursor/projects/workspace/uploads/WebexOne_2026_Template_0b88.pptx"
OUT = "/opt/cursor/artifacts/WebexOne2026_GenAI_Webex_Migration.pptx"
SVG_DIR = "/tmp/tpl/ppt/media"
ICON_CACHE = "/home/ubuntu/work/icons"
PHOTOS = {
    "cx_agent": "/opt/cursor/artifacts/assets/cx_agent.jpg",
    "receptionist": "/opt/cursor/artifacts/assets/ai_receptionist.jpg",
    "video_room": "/opt/cursor/artifacts/assets/video_room.jpg",
    "branch": "/opt/cursor/artifacts/assets/branch_team.jpg",
    "engineer_glasses": "/tmp/wx/ppt/media/image18.jpg",
    "datacenter": "/tmp/wx/ppt/media/image17.jpg",
    "noc": "/tmp/wx/ppt/media/image19.jpg",
    "lab": "/tmp/wx/ppt/media/image20.jpg",
}

INK = RGBColor(0x0A, 0x0A, 0x0A)
DARK = RGBColor(0x28, 0x28, 0x28)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
MID = RGBColor(0xB4, 0xB9, 0xC0)
LINE = RGBColor(0xD6, 0xD6, 0xD6)
PANEL = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CYAN = RGBColor(0x06, 0xA8, 0xFF)
BLUE = RGBColor(0x0A, 0x60, 0xFF)
VIOLET = RGBColor(0x8A, 0x3F, 0xFC)
PINK = RGBColor(0xFF, 0x00, 0x7F)
ORANGE = RGBColor(0xFF, 0x90, 0x00)
SKY = RGBColor(0x00, 0xC8, 0xFF)

HEAD = "Inter Semi Bold"
BODY = "Inter"
MONO = "Consolas"

ICONS = {
    "integration": 17, "automation": 18, "remote": 21, "idea": 22, "ai": 23,
    "cc_ai": 24, "people": 26, "transcript": 28, "apps": 29, "headset": 30,
    "chat": 32, "translate": 34, "audio": 35, "waveform": 41, "search_person": 44,
    "video": 46, "secure_cycle": 48, "link": 49, "privacy": 50, "triangle": 51,
    "lock": 53, "analytics": 55, "hidden": 56, "updown": 57, "mic": 58,
    "phone": 59, "docs": 60, "cx": 61, "stop": 62, "download": 63,
    "devices": 64, "graph": 65, "whiteboard": 66, "cloud": 67, "share": 33,
}

L_WALKIN, L_TITLE, L_AGENDA, L_TITLE_ONLY, L_TITLE_ONLY2 = 0, 1, 4, 6, 7
L_HALF_PHOTO, L_THIRD_PHOTO, L_TWO_THIRD_PHOTO = 17, 18, 19
L_SECTION_BODY_PHOTO = 21
L_QUOTE2 = 31
L_STATEMENT1, L_STATEMENT2, L_SECTION1, L_SECTION2, L_THANKS = 44, 45, 46, 47, 48


# ----------------------------------------------------------------- helpers
def hexs(c):
    return str(c)


def icon_png(name, color):
    os.makedirs(ICON_CACHE, exist_ok=True)
    num = ICONS[name]
    out = f"{ICON_CACHE}/{name}_{hexs(color)}.png"
    if not os.path.exists(out):
        svg = open(f"{SVG_DIR}/image{num}.svg").read()
        svg = re.sub(r"#282828", "#" + hexs(color), svg)
        cairosvg.svg2png(bytestring=svg.encode(), write_to=out, output_width=384)
    return out


def icon(slide, name, x, y, size, color=INK):
    return slide.shapes.add_picture(icon_png(name, color), Inches(x), Inches(y), Inches(size), Inches(size))


def _no_effects(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        eff = style.find(qn("a:effectRef"))
        if eff is not None:
            eff.set("idx", "0")


def set_fill(shape, fill):
    if fill is None:
        shape.fill.background()
    elif isinstance(fill, tuple) and fill[0] == "grad":
        _, c1, c2, angle = fill
        shape.fill.gradient()
        shape.fill.gradient_angle = angle
        stops = shape.fill.gradient_stops
        stops[0].position, stops[1].position = 0.0, 1.0
        stops[0].color.rgb, stops[1].color.rgb = c1, c2
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill


def set_line(shape, line, width=1.0, dash=None):
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
        if dash:
            ln = shape.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", dash)


def style_runs(p, runs, size, color, font, align, spacing=None):
    if align is not None:
        p.alignment = align
    if spacing:
        p.line_spacing = spacing
    for r in runs:
        if isinstance(r, str):
            r = (r, {})
        text, o = r
        run = p.add_run()
        run.text = text
        f = run.font
        f.size = Pt(o.get("size", size))
        f.name = o.get("font", font)
        f.bold = False
        f.italic = o.get("italic", False)
        c = o.get("color", color)
        if c is not None:
            f.color.rgb = c
        if o.get("strike"):
            run._r.get_or_add_rPr().set("strike", "sngStrike")


def fill_tf(tf, content, size=18, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            spacing=None, space_after=0, margins=(0, 0, 0, 0), wrap=True):
    tf.clear()
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        if child.tag in (qn("a:spAutoFit"), qn("a:normAutofit"), qn("a:noAutofit")):
            bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    if isinstance(content, (str, tuple)):
        content = [content]
    for i, para in enumerate(content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(para, dict):
            runs = para["runs"]
            style_runs(p, runs if isinstance(runs, list) else [runs], para.get("size", size),
                       para.get("color", color), para.get("font", font), para.get("align", align),
                       para.get("spacing", spacing))
            p.space_before = Pt(para.get("before", 0))
            p.space_after = Pt(para.get("after", space_after))
        else:
            style_runs(p, para if isinstance(para, list) else [para], size, color, font, align, spacing)
            p.space_after = Pt(space_after)
    return tf


def text(slide, x, y, w, h, content, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, content, **kw)
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12,
        content=None, line_width=1.0, dash=None, **kw):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(s, fill)
    set_line(s, line, line_width, dash)
    _no_effects(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    kw.setdefault("margins", (0.12, 0.06, 0.12, 0.06))
    kw.setdefault("anchor", MSO_ANCHOR.MIDDLE)
    kw.setdefault("align", PP_ALIGN.CENTER)
    fill_tf(s.text_frame, content if content is not None else "", **kw)
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=2.0, head=True, dash=None, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)
    if tail:
        he = etree.SubElement(ln, qn("a:headEnd"))
        he.set("type", "triangle"), he.set("w", "med"), he.set("len", "med")
    if head:
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle"), te.set("w", "med"), te.set("len", "med")
    return c


def circle(slide, x, y, d, fill, line=None, content=None, **kw):
    return box(slide, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL, radius=None, content=content,
               margins=(0, 0, 0, 0), **kw)


def icon_disc(slide, name, x, y, d, fill, icon_color=WHITE, ratio=0.56, line=None):
    circle(slide, x, y, d, fill, line=line)
    s = d * ratio
    icon(slide, name, x + (d - s) / 2, y + (d - s) / 2, s, icon_color)


def eyebrow(slide, label, color=BLUE, y=0.28):
    text(slide, 0.667, y, 11.5, 0.34, label.upper(), size=14, color=color, font=HEAD)


def set_title(slide, title):
    slide.shapes.title.text = title


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def grad(c1, c2, angle=0):
    return ("grad", c1, c2, angle)


def remove_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def new_slide(prs, layout_idx, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    if title is not None and s.shapes.title is not None:
        set_title(s, title)
    return s


def pipeline(slide, nodes, y, h=2.05, x0=0.667, total=12.0, gap=0.42, label_size=18, sub_size=14):
    n = len(nodes)
    w = (total - gap * (n - 1)) / n
    for i, nd in enumerate(nodes):
        x = x0 + i * (w + gap)
        hero = nd.get("hero")
        fill = nd.get("fill", PANEL)
        fg = WHITE if hero else INK
        sub_c = WHITE if hero else GREY
        box(slide, x, y, w, h, fill=fill, line=None if hero else LINE, radius=0.14)
        icon(slide, nd["icon"], x + w / 2 - 0.33, y + 0.2, 0.66, WHITE if hero else nd.get("icon_color", BLUE))
        text(slide, x + 0.1, y + 0.95, w - 0.2, 0.5, nd["label"], size=label_size, color=fg, font=HEAD,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)
        if nd.get("sub"):
            text(slide, x + 0.1, y + 1.45, w - 0.2, h - 1.5, nd["sub"], size=sub_size, color=sub_c,
                 align=PP_ALIGN.CENTER, spacing=0.95)
        if i < n - 1:
            arrow(slide, x + w + 0.06, y + h / 2, x + w + gap - 0.06, y + h / 2, color=GREY, width=2.25)
    return w


# ----------------------------------------------------------------- deck
def build():
    prs = Presentation(TEMPLATE)
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)

    # 0. Walk-in
    s = new_slide(prs, L_WALKIN)
    notes(s, """
Walk-in slide. Leave this up while the room fills.
Optional: play the session title and your name on the confidence monitor.
""")

    # 1. Title
    s = new_slide(prs, L_TITLE)
    set_title(s, "Legacy to AI-First:\nGenAI-Powered Migrations")
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 13:
            ph.text = "Five on-prem workloads to Webex. One repeatable pattern. AI at both ends."
        elif idx == 11:
            ph.text = "Speaker Name\nTitle, Company"
        elif idx == 12:
            ph.text = "WebexOne 2026  |  Technical Session"
    remove_empty_placeholders(s)
    notes(s, """
Welcome. Today is about moving five entrenched on-prem Cisco workloads - UCCX, CUCM Hunt Pilots, Unity Connection, CUCM dial plans, and Collaboration Endpoints - into their cloud-native Webex counterparts.
The twist: AI shows up twice. First as the engineering copilot that parses and translates legacy configuration. Then as the destination - Webex AI Agent, AI Receptionist, Voicemail AI Transcription, and the Webex AI Assistant.
Attendees leave with prompt-driven frameworks and blueprints they can use on Monday.
""")

    # 2. Agenda
    s = new_slide(prs, L_AGENDA, "Agenda")
    items = ["Why migrate now", "AI works twice", "The five-step method", "Five plays, one pattern",
             "Guardrails & takeaways"]
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx in (18, 20):
            ph._element.getparent().remove(ph._element)
    for i, item in enumerate(items):
        y = 0.62 + i * 0.8
        text(s, 5.229, y, 0.9, 0.5, f"0{i + 1}", size=24, color=BLUE, font=HEAD)
        text(s, 6.4635, y, 6.2, 0.5, item, size=24, color=INK, font=HEAD)
    remove_empty_placeholders(s)
    notes(s, """
Five stops:
1. Why the legacy stack is holding us back.
2. The two roles AI plays - during the migration and after it.
3. The five-step method: Extract, Sanitize, Analyze, Transform, Validate.
4. Five migration plays: UCCX, Hunt Pilots, Unity, Dial Plan, Devices.
5. Guardrails that keep AI honest, and the takeaways.
""")

    # 3. Statement
    s = new_slide(prs, L_STATEMENT2, "Don't just migrate.\nTransform.")
    notes(s, """
The core message of this session. A like-for-like lift of a 15-year-old IVR tree or a tangle of Hunt Pilots just moves technical debt to the cloud.
The migration is the moment to modernize the customer and employee experience - and to land on Cisco's native AI portfolio.
""")

    # 4. The legacy tax
    s = new_slide(prs, L_TITLE_ONLY, "The legacy tax")
    eyebrow(s, "01 · Why migrate now")
    tiles = [
        ("UCCX", "Java-serialized AEF scripts", "headset", BLUE),
        ("Hunt Pilots", "Dead pilots, tangled lists", "people", CYAN),
        ("Unity Connection", "Deep touch-tone trees", "phone", PINK),
        ("CUCM Dial Plan", "CSS & partition sprawl", "graph", ORANGE),
        ("CE Endpoints", "Aging macros & xAPI", "devices", VIOLET),
    ]
    w, gap = 2.2, 0.25
    for i, (name, pain, ic, col) in enumerate(tiles):
        x = 0.667 + i * (w + gap)
        box(s, x, 1.55, w, 3.35, fill=PANEL, line=None, radius=0.16)
        box(s, x, 1.55, w, 0.1, fill=col, radius=None, shape=MSO_SHAPE.RECTANGLE)
        icon_disc(s, ic, x + w / 2 - 0.55, 1.95, 1.1, col)
        text(s, x + 0.05, 3.25, w - 0.1, 0.5, name, size=19, font=HEAD, align=PP_ALIGN.CENTER, wrap=False)
        text(s, x + 0.15, 3.85, w - 0.3, 0.9, pain, size=16, color=GREY, align=PP_ALIGN.CENTER, spacing=0.95)
    text(s, 0.667, 5.35, 12.0, 0.9, [[("Opaque.   ", {"color": BLUE}), ("Manual.   ", {"color": PINK}),
                                      ("Error-prone.", {"color": ORANGE})]],
         size=44, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
Traditional architectures span five rigid domains. Extracting routing logic from them and adapting it to cloud formats is slow and error-prone:
- UCCX: Application Editor Format (.aef) scripts are Java-serialized - standard text tools can't read them.
- CUCM Hunt Pilots: years of "dead" pilots, hunt lists and line groups nobody owns.
- Unity Connection: multi-level DTMF call-handler trees, exported as massive CSVs.
- CUCM dial plan: partitions, calling search spaces, route patterns and shared lines - traditionally managed through SQL or BAT.
- CE endpoints: JavaScript macros and UI Extensions written against CE 9.x.
Manual porting stifles agility and delays cloud adoption. That's the tax.
""")

    # 5. AI works twice
    s = new_slide(prs, L_TITLE_ONLY, "AI works twice")
    eyebrow(s, "02 · Two roles for AI")
    box(s, 0.667, 1.5, 4.7, 4.95, fill=PANEL, line=None, radius=0.2)
    text(s, 0.95, 1.7, 4.2, 0.4, "DURING THE MIGRATION", size=15, color=BLUE, font=HEAD)
    text(s, 0.95, 2.05, 4.2, 0.55, "AI as your copilot", size=26, font=HEAD)
    left = [("docs", "Parse legacy configs"), ("link", "Map to cloud objects"),
            ("download", "Draft CSV & JSON"), ("analytics", "Generate test plans")]
    for i, (ic, lbl) in enumerate(left):
        y = 2.85 + i * 0.86
        box(s, 0.95, y, 4.15, 0.7, fill=WHITE, line=LINE, radius=0.12)
        icon(s, ic, 1.1, y + 0.12, 0.46, BLUE)
        text(s, 1.75, y, 3.3, 0.7, lbl, size=19, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    icon_disc(s, "cloud", 5.87, 2.95, 1.6, grad(SKY, BLUE, 45), ratio=0.55)
    text(s, 5.47, 4.65, 2.4, 0.5, "Webex Cloud", size=20, font=HEAD, align=PP_ALIGN.CENTER)
    arrow(s, 5.42, 3.75, 5.82, 3.75, color=GREY, width=2.5)
    arrow(s, 7.52, 3.75, 7.92, 3.75, color=GREY, width=2.5)
    box(s, 7.967, 1.5, 4.7, 4.95, fill=grad(BLUE, PINK, 315), line=None, radius=0.2)
    text(s, 8.25, 1.7, 4.2, 0.4, "AFTER THE MIGRATION", size=15, color=WHITE, font=HEAD)
    text(s, 8.25, 2.05, 4.2, 0.55, "AI as your product", size=26, color=WHITE, font=HEAD)
    right = [("cc_ai", "Webex AI Agent"), ("mic", "AI Receptionist"),
             ("transcript", "Voicemail Transcription"), ("devices", "AI Assistant on RoomOS")]
    for i, (ic, lbl) in enumerate(right):
        y = 2.85 + i * 0.86
        box(s, 8.25, y, 4.15, 0.7, fill=WHITE, line=None, radius=0.12)
        icon(s, ic, 8.4, y + 0.12, 0.46, PINK)
        text(s, 9.05, y, 3.3, 0.7, lbl, size=19, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
A central theme: AI plays two distinct roles.
Role 1 - AI during the transition: LLMs analyze configuration exports, summarize routing behavior, discover dependencies, detect conflicts, recommend mappings and generate draft artifacts (CSV, JSON, code, test plans).
Role 2 - AI enabled by the transition: once you land in the cloud you unlock Webex AI Agent, AI Receptionist, Voicemail AI Transcription, and Webex AI Assistant / GenAI on RoomOS.
This dual approach is how organizations get to an "AI-first" operating state efficiently.
""")

    # 6. Section - method
    s = new_slide(prs, L_SECTION1, "01\u000bThe Method")
    notes(s, "Section: the repeatable five-step method we apply to every workload.")

    # 7. Five-step method
    s = new_slide(prs, L_TITLE_ONLY, "One method. Every workload.")
    eyebrow(s, "The five-step method")
    steps = [("Extract", "Pull source data", "download", CYAN),
             ("Sanitize", "Strip PII & secrets", "privacy", BLUE),
             ("Analyze", "Recover intent", "ai", VIOLET),
             ("Transform", "Map to Webex", "automation", PINK),
             ("Validate", "Test every artifact", "secure_cycle", ORANGE)]
    step, w = 2.37, 2.55
    for i, (name, desc, ic, col) in enumerate(steps):
        x = 0.667 + i * step
        icon_disc(s, ic, x + w / 2 - 0.5, 1.55, 1.0, col)
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        c = box(s, x, 2.8, w, 1.2, fill=col, shape=shp, radius=None,
                content=[[(name, {"font": HEAD})]], size=22, color=WHITE, margins=(0.3 if i else 0.05, 0, 0.1, 0),
                wrap=False)
        c.adjustments[0] = 0.32
        text(s, x + 0.1, 4.15, w - 0.3, 0.45, desc, size=18, color=GREY, align=PP_ALIGN.CENTER)
    y_loop = 5.1
    arrow(s, 10.55, 4.65, 10.55, y_loop, color=MID, width=2, head=False)
    arrow(s, 10.55, y_loop, 5.9, y_loop, color=MID, width=2, head=False)
    arrow(s, 5.9, y_loop, 5.9, 4.68, color=MID, width=2)
    text(s, 6.3, y_loop - 0.02, 3.9, 0.4, "Iterate until it passes", size=15, color=GREY,
         align=PP_ALIGN.CENTER)
    box(s, 2.67, 5.75, 8.0, 0.7, fill=INK, radius=0.35,
        content=[[("Human in the loop:  ", {"color": SKY}), ("AI drafts. Engineers approve.", {})]],
        size=20, color=WHITE, font=HEAD)
    notes(s, """
Every one of the five plays follows the same pattern:
1. Extract - pull supported source data reliably (BAT exports, CallHandlerDataDump, AEF XML, device inventory).
2. Sanitize - strip sensitive information before anything touches an LLM.
3. Analyze - recover behavior, relationships and intent, not just objects.
4. Transform - turn findings into target-state recommendations and draft artifacts (Webex Calling CSVs, WxCC flow JSON, RoomOS macros).
5. Validate - every artifact is tested before deployment. When it fails, loop back to analysis.
Humans stay accountable at every gate.
""")

    # 8. Migration map
    s = new_slide(prs, L_TITLE_ONLY, "Five workloads. Five AI destinations.")
    eyebrow(s, "The migration map")
    cols = [(0.667, 2.55, "LEGACY"), (3.45, 3.25, "HOW IT MOVES"), (6.93, 2.6, "WEBEX TARGET"),
            (9.77, 2.9, "AI PAYOFF")]
    for x, w_, lbl in cols:
        text(s, x, 1.42, w_, 0.3, lbl, size=13, color=GREY, font=HEAD, align=PP_ALIGN.CENTER)
    rows = [
        ("UCCX", "Auto-extract CSQs & skills", "Webex Contact Center", "Webex AI Agent", BLUE),
        ("Hunt Pilots", "Control Hub bulk provisioning", "Customer Assist", "AI summaries + sentiment", CYAN),
        ("Unity Connection", "CallHandlerDataDump \u2192 CSV", "Auto Attendants", "AI Receptionist", PINK),
        ("CUCM Dial Plan", "Migration wizard + E.164", "Webex Calling", "Voicemail AI Transcription", ORANGE),
        ("CE Endpoints", "Webex Device Connector", "RoomOS Workspaces", "AI Assistant + GenAI macros", VIOLET),
    ]
    for i, (leg, how, tgt, ai, col) in enumerate(rows):
        y = 1.82 + i * 0.95
        box(s, 0.667, y, 2.55, 0.78, fill=PANEL, radius=0.12, content=leg, size=19, font=HEAD)
        box(s, 0.667, y, 0.09, 0.78, fill=col, shape=MSO_SHAPE.RECTANGLE, radius=None)
        arrow(s, 3.3, y + 0.52, 6.85, y + 0.52, color=MID, width=1.75)
        text(s, 3.45, y + 0.02, 3.25, 0.42, how, size=15, color=GREY, align=PP_ALIGN.CENTER)
        box(s, 6.93, y, 2.6, 0.78, fill=col, radius=0.12, content=tgt, size=18, color=WHITE, font=HEAD)
        arrow(s, 9.57, y + 0.39, 9.73, y + 0.39, color=GREY, width=1.75)
        box(s, 9.77, y, 2.9, 0.78, fill=grad(BLUE, PINK, 0), radius=0.12, content=ai, size=17, color=WHITE,
            font=HEAD)
    notes(s, """
The migration map - parity preservation plus immediate access to Cisco's native AI portfolio.
- UCCX -> Webex Contact Center. Automated extraction of CSQs, agents, skills and prompts into WxCC entry points and queues. Destination: Webex AI Agent, a conversational omnichannel digital front door replacing legacy IVR.
- CUCM Hunt Pilot -> Webex Customer Assist. Control Hub bulk provisioning maps Hunt Pilots/Lists to Customer Assist queues; Local Gateway handles hybrid routing. Destination: AI Agent summaries - real-time and post-call summaries, sentiment scoring (2025-26 roadmap).
- Unity Connection -> Webex Calling Auto Attendants. CallHandlerDataDump (via Informix ODBC) extracts zip/CSVs for Control Hub bulk import. Destination: Webex AI Receptionist - 24/7 NLP intent routing replacing DTMF menus.
- CUCM Dial Plan -> Webex Calling. Control Hub migration wizards map route patterns and DNs using E.164 normalization and flat extensions. Destination: Voicemail AI Transcription in the Webex App.
- CE devices -> Webex Cloud (RoomOS). Webex Device Connector transfers and centralizes management as Workspaces in Control Hub. Destination: Webex AI Assistant, real-time captioning, and LLM-assisted macro translation.
""")

    # 9. Quote
    s = new_slide(prs, L_QUOTE2, "AI drafts.\nEngineers decide.")
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 12:
            ph.text = "Your copilot - never the authority."
    remove_empty_placeholders(s)
    notes(s, """
AI is a copilot, not an authority:
- Summarize large configuration sets to accelerate analysis.
- Identify patterns and map relationships.
- Draft technical documentation and test cases.
- But validate every generated configuration - engineering accountability stays with us.
- Protect production data: never upload it without authorization.
""")

    # 10. Prompt architecture
    s = new_slide(prs, L_TITLE_ONLY, "Prompt like an architect")
    eyebrow(s, "Prompt architecture")
    layers = [("Role", "Act as a UCCX migration analyst", CYAN),
              ("Source", "Sanitized export + product version", BLUE),
              ("Constraints", "Schema, privacy, no invented APIs", VIOLET),
              ("Output", "JSON / CSV + mapping table", PINK),
              ("Uncertainty", "Flag REQUIRES_VALIDATION", ORANGE)]
    for i, (name, ex, col) in enumerate(layers):
        y = 1.55 + i * 0.98
        box(s, 0.667, y, 5.6, 0.84, fill=PANEL, radius=0.12)
        box(s, 0.667, y, 0.84, 0.84, fill=col, radius=0.12, content=str(i + 1), size=24, color=WHITE, font=HEAD)
        text(s, 1.7, y + 0.06, 4.4, 0.4, name, size=20, font=HEAD)
        text(s, 1.7, y + 0.44, 4.5, 0.35, ex, size=15, color=GREY)
    box(s, 6.75, 1.55, 5.92, 4.76, fill=INK, radius=0.18)
    for j, c in enumerate([PINK, ORANGE, SKY]):
        circle(s, 7.0 + j * 0.28, 1.75, 0.16, c)
    text(s, 8.0, 1.66, 3.5, 0.35, "migration_prompt.txt", size=13, color=MID, font=MONO)
    code = [("ROLE    ", "UCCX \u2192 WxCC migration analyst"),
            ("SOURCE  ", "sanitized .aef XML, UCCX 12.5"),
            ("GOAL    ", "Flow Designer-ready JSON"),
            ("RULES   ", "documented nodes only"),
            ("        ", "never invent APIs or fields"),
            ("OUTPUT  ", "flow.json + mapping.csv"),
            ("FLAG    ", "REQUIRES_VALIDATION if unsure")]
    text(s, 7.0, 2.2, 5.5, 3.9, [[(k, {"color": SKY}), (v, {"color": WHITE})] for k, v in code],
         size=17, font=MONO, space_after=7)
    notes(s, """
Prompt architecture for reliable results - five layers:
1. Define role, source and goal - establish context so the model reasons like a specialist.
2. State the source - sanitized configuration, with product and version.
3. State schema and privacy constraints - protect data while keeping structure; use only documented APIs and fields.
4. Specify the output format - machine-readable JSON/CSV plus a human-readable mapping table and reasoning artifacts, so you can trace logic.
5. Specify uncertainty handling - anything the model is unsure about is labeled REQUIRES_VALIDATION instead of silently guessed.
The example on the right is the skeleton we use for UCCX; swap the role/source/goal for each play.
""")

    # 11. Sanitize before you send
    s = new_slide(prs, L_TITLE_ONLY, "Sanitize before you send")
    eyebrow(s, "Data guardrails")
    flow = [("docs", "Raw export", GREY), ("privacy", "Tokenize", PINK), ("ai", "LLM analysis", BLUE),
            ("secure_cycle", "Re-hydrate locally", CYAN)]
    for i, (ic, lbl, col) in enumerate(flow):
        x = 0.9 + i * 3.1
        icon_disc(s, ic, x + 0.55, 1.5, 1.2, col)
        text(s, x - 0.2, 2.8, 2.7, 0.45, lbl, size=19, font=HEAD, align=PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, x + 1.9, 2.1, x + 3.5, 2.1, color=MID, width=2.25)
    box(s, 0.667, 1.35, 12.0, 2.05, fill=None, line=MID, dash="dash", radius=0.2)
    rows = [("Phone", "+1 703 555 1212", "+1XXXXXXXX01"),
            ("User", "jane.smith@example.com", "USER_014"),
            ("Site", "HQ-Route-List", "SITE_A_ROUTE_LIST")]
    for i, (k, src, dst) in enumerate(rows):
        y = 3.85 + i * 0.78
        text(s, 0.667, y, 1.2, 0.6, k, size=17, color=GREY, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
        box(s, 1.9, y, 4.3, 0.6, fill=PANEL, radius=0.1, content=src, size=17, font=MONO, color=DARK)
        arrow(s, 6.3, y + 0.3, 7.25, y + 0.3, color=PINK, width=2.25)
        box(s, 7.35, y, 4.3, 0.6, fill=WHITE, line=BLUE, line_width=1.5, radius=0.1, content=dst, size=17,
            font=MONO, color=BLUE)
    icon(s, "lock", 11.95, 4.0, 0.6, INK)
    text(s, 11.75, 4.65, 1.0, 1.2, "Key stays local", size=14, color=GREY, align=PP_ALIGN.CENTER)
    notes(s, """
Data guardrails. Before any configuration goes to an LLM:
- Tokenize phone numbers, user identities, site names and anything sensitive - consistently, so relationships survive (the same site always becomes SITE_A).
- Keep the mapping key inside your network. The model only ever sees tokens.
- After analysis, re-hydrate the output locally.
- Use an approved enterprise model/tenant; never upload production data without authorization.
Bonus: consistent tokens also make LLM outputs easier to diff and test.
""")

    # 12. Section - plays
    s = new_slide(prs, L_SECTION1, "02\u000bFive Plays")
    notes(s, "Section: five migration plays - UCCX, Hunt Pilots, Unity, Dial Plan, Devices.")

    # 13. Play 1 - UCCX pipeline
    s = new_slide(prs, L_TITLE_ONLY, "Decode the AEF. Rebuild in Flow Designer.")
    eyebrow(s, "Play 1 · UCCX \u2192 Webex Contact Center", BLUE)
    pipeline(s, [
        {"icon": "docs", "label": ".aef script", "sub": "Java-serialized XML", "icon_color": GREY},
        {"icon": "ai", "label": "LLM analyst", "sub": "Claude · ChatGPT · Gemini", "hero": True,
         "fill": grad(SKY, BLUE, 45)},
        {"icon": "download", "label": "Flow JSON", "sub": "Import-ready"},
        {"icon": "graph", "label": "Flow Designer", "sub": "Webex Contact Center"},
        {"icon": "cc_ai", "label": "AI Agent", "sub": "Webex digital front door", "hero": True,
         "fill": grad(BLUE, PINK, 45)},
    ], y=1.55)
    text(s, 0.667, 3.95, 6, 0.35, "WHAT MOVES", size=13, color=GREY, font=HEAD)
    chips = ["CSQs \u2192 Queues", "Skills", "Agents & teams", "Prompts", "Wrap-up codes", "Entry points"]
    x = 0.667
    widths = [2.35, 1.2, 2.25, 1.5, 2.2, 2.0]
    for c, cw in zip(chips, widths):
        box(s, x, 4.35, cw, 0.62, fill=WHITE, line=BLUE, line_width=1.5, radius=0.31, content=c, size=16,
            color=BLUE, font=HEAD, wrap=False, margins=(0.04, 0, 0.04, 0))
        x += cw + 0.1
    text(s, 0.667, 5.45, 12.0, 0.8, [[("No manual rebuilds.  ", {"color": INK}),
                                      ("Review. Test. Import.", {"color": BLUE})]],
         size=34, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
Play 1 - UCCX to Webex Contact Center.
- UCCX Application Editor Format (.aef) files are Java-serialized and hard to parse with standard text tools.
- Engineers copy the encoded XML into an LLM (Claude, ChatGPT, Gemini) and prompt it to act as a UCCX analyst.
- The LLM identifies legacy nodes (Select Resource, Menu, etc.), recovers entry points and triggers, menu options and retry logic, business-hours and holiday routing, CSQ/skill dependencies, and error/timeout treatment.
- It generates Webex Contact Center-compatible JSON for import into Flow Designer.
- In parallel, the automated migration tooling extracts CSQs, agents, skills, prompts and wrap-up codes into WxCC entry points and queue objects - bridging static rule-based routing and dynamic cloud entry points.
End state: deploy the Webex AI Agent as the omnichannel digital front door.
""")

    # 14. Play 1 - intent mapping
    s = new_slide(prs, L_TITLE_ONLY, "Every UCCX step has a new home")
    eyebrow(s, "Play 1 · Intent mapping", BLUE)
    for x, w_, lbl in [(0.667, 3.3, "UCCX STEP"), (4.1, 3.9, "PURPOSE"), (8.2, 4.47, "WEBEX CONTACT CENTER")]:
        text(s, x, 1.33, w_, 0.3, lbl, size=13, color=GREY, font=HEAD, align=PP_ALIGN.CENTER)
    maps = [("Menu", "Capture choice", "Menu \u2192 AI intent"),
            ("Select Resource", "Route to agent", "Queue + skill match"),
            ("Day of Week", "Apply schedule", "Business-hours check"),
            ("Get Contact Info", "Read attributes", "Flow variables"),
            ("Database step", "Fetch data", "HTTP request / API"),
            ("Prompt", "Inform caller", "TTS or AI response"),
            ("Exception branch", "Handle failure", "Error path")]
    for i, (a, p, b) in enumerate(maps):
        y = 1.7 + i * 0.68
        box(s, 0.667, y, 3.3, 0.56, fill=PANEL, radius=0.1, content=a, size=18, font=HEAD)
        arrow(s, 4.05, y + 0.28, 8.12, y + 0.28, color=LINE, width=1.5)
        box(s, 4.75, y + 0.06, 2.6, 0.44, fill=WHITE, radius=0.22, content=p, size=16, color=GREY,
            wrap=False, margins=(0.02, 0, 0.02, 0))
        box(s, 8.2, y, 4.47, 0.56, fill=BLUE if i else grad(BLUE, PINK, 0), radius=0.1, content=b, size=18,
            color=WHITE, font=HEAD)
    notes(s, """
The LLM's job is to recover the business purpose of each step, then propose a WxCC treatment:
- Menu: capture caller choice -> Menu activity, or better, an AI Agent intent.
- Select Resource: route to agent -> queue with skill matching.
- Day of Week: apply schedule -> business-hours logic.
- Get Contact Info: read attributes -> flow variables.
- Database step: retrieve data -> HTTP request / API integration.
- Prompt: give information -> TTS or AI response.
- Exception branch: handle failure -> error path.
Target design principles: preserve valid experiences, simplify by removing unnecessary menus, transform rigid interactions with AI, and provide alternatives - not silent decisions.
""")

    # 15. Play 1 - AI Agent
    s = new_slide(prs, L_HALF_PHOTO, "Retire the IVR.\nOpen a front door.")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 18:
            ph.insert_picture(PHOTOS["cx_agent"])
    eyebrow(s, "Play 1 · Webex AI Agent", BLUE, y=0.25)
    box(s, 0.667, 2.0, 5.3, 0.7, fill=PANEL, radius=0.35,
        content=[[("\u201cPress 1 for billing. Press 2 for\u2026\u201d", {"strike": True})]], size=18, color=GREY)
    box(s, 0.667, 2.85, 5.3, 0.8, fill=grad(BLUE, PINK, 0), radius=0.4,
        content="\u201cHi! How can I help today?\u201d", size=22, color=WHITE, font=HEAD)
    feats = [("ai", "Autonomous + scripted modes"), ("apps", "Acts in Epic & Salesforce"),
             ("translate", "Empathetic, multilingual voice"), ("transcript", "Hands off with a summary")]
    for i, (ic, lbl) in enumerate(feats):
        y = 4.05 + i * 0.66
        icon(s, ic, 0.72, y, 0.48, BLUE)
        text(s, 1.4, y - 0.02, 4.8, 0.52, lbl, size=19, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
Cisco is phasing out touch-tone and rigid scripted voice menus - including Nuance-based IVR - in favor of conversational AI.
Webex AI Agent is the "digital front door":
- Runs in autonomous (LLM-driven) and scripted modes.
- Integrates with back-end systems like Epic EHR and Salesforce for appointment scheduling, triage and other tasks.
- Uses ElevenLabs voices (from mid-2025) for empathetic, multilingual text-to-speech.
- AI Agent Studio in Control Hub ingests documents using parent-child chunking to generate grounded answers.
- On escalation, Real-Time Assist hands a generated interaction summary to the human agent.
Implementation: extracted UCCX logic and Nuance static database dips move into Webex Connect Flow Builder as intent-based natural-language workflows.
""")

    # 16. Play 2 - right-size
    s = new_slide(prs, L_TITLE_ONLY, "Right-size every Hunt Pilot")
    eyebrow(s, "Play 2 · CUCM Hunt Pilots \u2192 Webex Customer Assist", CYAN)
    text(s, 0.667, 1.55, 1.9, 0.55, "TODAY", size=14, color=GREY, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    chain = ["Hunt Pilot", "Hunt List", "Line Group", "Members"]
    for i, c in enumerate(chain):
        x = 1.75 + i * 2.75
        box(s, x, 1.55, 2.25, 0.55, fill=PANEL, radius=0.27, content=c, size=17, font=HEAD, color=DARK)
        if i < 3:
            arrow(s, x + 2.3, 1.825, x + 2.7, 1.825, color=MID, width=2)
    stairs = [("Hunt Group", "Simple ringing", 1.35, MID, INK),
              ("Call Queue", "Basic queuing", 1.95, CYAN, WHITE),
              ("Customer Assist", "Frontline teams", 2.6, BLUE, WHITE),
              ("Contact Center", "Advanced CX", 3.25, grad(BLUE, PINK, 45), WHITE)]
    base = 6.3
    for i, (dest, need, hgt, col, fg) in enumerate(stairs):
        x = 0.667 + i * 3.05
        box(s, x, base - hgt, 2.85, hgt, fill=col, radius=0.16)
        text(s, x + 0.2, base - hgt + 0.2, 2.5, 0.45, dest, size=21, color=fg, font=HEAD)
        text(s, x + 0.2, base - hgt + 0.65, 2.5, 0.4, need, size=16, color=fg)
    box(s, 6.95, 3.2, 1.8, 0.42, fill=ORANGE, radius=0.21, content="Sweet spot", size=15, color=WHITE,
        font=HEAD)
    arrow(s, 0.667, 6.5, 12.6, 6.5, color=MID, width=1.5)
    text(s, 9.2, 6.52, 3.4, 0.35, "More capability", size=14, color=GREY, align=PP_ALIGN.RIGHT)
    notes(s, """
Play 2 - CUCM Hunt Pilots.
Review the call-flow chain end to end: Hunt Pilot (entry point) -> Hunt List (ordered destinations) -> Line Groups (member grouping, distribution algorithm) -> member DNs. Also check CSS/partitions, no-answer timers, overflow and schedules.
Then right-size the destination:
- Simple ringing -> Webex Calling Hunt Group.
- Basic queuing -> Webex Calling Call Queue.
- Frontline staff / branch teams -> Webex Customer Assist - the sweet spot for organizations that don't need a full contact center: advanced queuing, skill-based routing, supervisor dashboards in the Webex App. Requires the Webex Calling Professional license.
- Advanced routing and omnichannel -> Webex Contact Center.
""")

    # 17. Play 2 - clean first
    s = new_slide(prs, L_THIRD_PHOTO, "Clean first. Then move.")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 18:
            ph.insert_picture(PHOTOS["branch"])
    eyebrow(s, "Play 2 · Rationalize with GenAI", CYAN, y=0.25)
    flow = ["BAT extract", "Rebuild links", "Flag dead", "Classify", "Ask owners"]
    for i, f in enumerate(flow):
        x = 0.667 + i * 1.5
        circle(s, x + 0.35, 1.55, 0.7, CYAN if i != 2 else PINK, content=str(i + 1), size=20, color=WHITE,
               font=HEAD)
        text(s, x - 0.05, 2.35, 1.5, 0.6, f, size=15, font=HEAD, align=PP_ALIGN.CENTER, spacing=0.9)
        if i < 4:
            arrow(s, x + 1.1, 1.9, x + 1.8, 1.9, color=MID, width=1.75)
    text(s, 0.667, 3.2, 7.0, 0.35, "CLASSIFY EVERY PILOT", size=13, color=GREY, font=HEAD)
    cls = [("Preserve", BLUE), ("Simplify", CYAN), ("Modernize", PINK), ("Retire", DARK), ("Need info", ORANGE)]
    for i, (c, col) in enumerate(cls):
        box(s, 0.667 + i * 1.47, 3.6, 1.37, 0.6, fill=col, radius=0.3, content=c, size=15, color=WHITE,
            font=HEAD, wrap=False, margins=(0.02, 0, 0.02, 0))
    box(s, 0.667, 4.6, 7.2, 1.75, fill=grad(CYAN, BLUE, 0), radius=0.18)
    icon(s, "analytics", 0.95, 4.85, 0.7, WHITE)
    text(s, 1.85, 4.8, 5.8, 0.5, "The AI payoff", size=22, color=WHITE, font=HEAD)
    text(s, 1.85, 5.3, 5.9, 0.9, ["Call summaries  ·  Sentiment scoring", "Supervisor dashboards in Webex App"],
         size=17, color=WHITE, spacing=1.05)
    notes(s, """
Rationalizing Hunt Pilots with GenAI:
1. Upload a sanitized CUCM BAT extract.
2. Have the LLM reconstruct Hunt Pilot -> Hunt List -> Line Group -> member relationships.
3. Flag empty, inconsistent or inactive ("dead") groups - clean these up in CUCM before transition so you don't import redundant data.
4. Classify each group: Preserve, Simplify, Modernize, Retire, or Need more information. (Example: a 5-member help desk might be preserved; a 20-member service team simplified; a pilot inactive for six months retired.)
5. Generate business-owner questions to validate requirements.
Payoff on Customer Assist: real-time and post-call AI summaries and sentiment scoring (2025-26 roadmap), plus supervisor dashboards.
Hybrid safety: Control Hub bulk provisioning maps Hunt Pilots/Lists to Customer Assist queues, while Local Gateway handles hybrid routing during the phased move.
""")

    # 18. Play 3 - Unity pipeline
    s = new_slide(prs, L_TITLE_ONLY, "Unity handlers \u2192 AI Receptionist")
    eyebrow(s, "Play 3 · Unity Connection \u2192 Webex Calling", PINK)
    pipeline(s, [
        {"icon": "phone", "label": "Unity", "sub": "Connection call handlers", "icon_color": GREY},
        {"icon": "docs", "label": "Data Dump", "sub": "ZIP of CSVs", "icon_color": PINK},
        {"icon": "ai", "label": "LLM mapping", "sub": "Trees · orphans · CSV", "hero": True,
         "fill": grad(PINK, VIOLET, 45)},
        {"icon": "download", "label": "Control Hub", "sub": "Bulk-import Auto Attendants", "icon_color": PINK},
        {"icon": "mic", "label": "AI Receptionist", "sub": "NLP intent routing", "hero": True,
         "fill": grad(BLUE, PINK, 45)},
    ], y=1.55, sub_size=14)
    box(s, 0.667, 4.1, 5.85, 1.95, fill=PANEL, radius=0.18)
    text(s, 0.95, 4.3, 5.3, 0.35, "PREREQUISITES", size=13, color=GREY, font=HEAD)
    text(s, 0.95, 4.7, 5.4, 1.2, ["32-bit Informix ODBC driver", "Remote Database Proxy enabled"],
         size=20, font=HEAD, space_after=8)
    box(s, 6.82, 4.1, 5.85, 1.95, fill=WHITE, line=ORANGE, line_width=2, radius=0.18)
    icon(s, "audio", 7.1, 4.35, 0.62, ORANGE)
    text(s, 7.95, 4.3, 4.5, 0.35, "HEADS UP", size=13, color=ORANGE, font=HEAD)
    text(s, 7.95, 4.7, 4.55, 1.2, ["Greetings (.wav) don't migrate.", "Re-upload them manually."], size=20,
         font=HEAD, space_after=8)
    notes(s, """
Play 3 - Unity Connection call handlers.
- Extraction: the CallHandlerDataDump Windows utility. It requires a 32-bit Informix ODBC driver and the Remote Database Proxy service enabled on Unity Connection. Output is a zip of CSVs.
- The CSVs are massive and complex. The LLM maps the multi-level IVR tree, identifies abandoned/unreferenced call handlers, routing loops and missing destinations, and translates the active data into the exact CSV format Control Hub expects for bulk-provisioning Webex Calling Auto Attendants.
- Touchpoints to preserve: handler data, caller-input key mappings, transfer rules, standard/closed/holiday greetings, schedules, interview and directory handlers, extension relationships.
- Watch-out: audio greetings (.wav) are not parsed automatically - plan to re-upload them into Webex Calling.
End state: upgrade to the NLP-driven Webex AI Receptionist.
""")

    # 19. Play 3 - key presses to intents
    s = new_slide(prs, L_TITLE_ONLY, "From key presses to intents")
    eyebrow(s, "Play 3 · Redesign the experience", PINK)
    box(s, 1.95, 1.55, 2.6, 0.62, fill=DARK, radius=0.12, content="Main greeting", size=17, color=WHITE,
        font=HEAD)
    kids = ["1  Sales", "2  Support", "3  Billing", "0  Operator"]
    for i, k in enumerate(kids):
        x = 0.667 + i * 1.33
        arrow(s, 3.25, 2.17, x + 0.6, 2.75, color=MID, width=1.25, head=False)
        box(s, x, 2.75, 1.2, 0.55, fill=PANEL, radius=0.1, content=k, size=14, font=HEAD, color=DARK,
            wrap=False, margins=(0.02, 0, 0.02, 0))
    for i, k in enumerate(["1  Hardware", "2  Software", "3  Other"]):
        y = 3.6 + i * 0.62
        arrow(s, 2.6, 3.3, 2.6, y + 0.27, color=MID, width=1.25, head=False)
        box(s, 2.0, y, 1.2, 0.5, fill=PANEL, radius=0.1, content=k, size=13, color=GREY, wrap=False,
            margins=(0.02, 0, 0.02, 0))
    text(s, 0.667, 5.65, 5.2, 0.6, "Rigid. Deep. Touch-tone.", size=24, color=GREY, font=HEAD,
         align=PP_ALIGN.CENTER)
    box(s, 5.95, 3.1, 1.1, 0.9, fill=grad(BLUE, PINK, 0), shape=MSO_SHAPE.RIGHT_ARROW, radius=None)
    bub = box(s, 7.45, 1.55, 5.2, 1.0, fill=grad(BLUE, PINK, 0), shape=MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
              radius=None, content="\u201cMy card was charged twice.\u201d", size=22, color=WHITE, font=HEAD)
    bub.adjustments[0], bub.adjustments[1] = -0.3, 0.8
    box(s, 7.45, 3.0, 5.2, 0.7, fill=WHITE, line=PINK, line_width=2, radius=0.35,
        content=[[("Intent  ", {"color": GREY}), ("Billing dispute", {})]], size=20, font=HEAD)
    arrow(s, 10.05, 3.72, 10.05, 4.08, color=PINK, width=2)
    box(s, 7.45, 4.1, 5.2, 0.7, fill=PINK, radius=0.35,
        content=[[("Route  ", {"color": WHITE}), ("Billing team", {"color": WHITE})]], size=20, font=HEAD)
    box(s, 7.45, 4.95, 5.2, 0.6, fill=PANEL, radius=0.3,
        content=[[("Low confidence?  ", {"color": GREY}), ("\u2192 Operator", {"color": INK})]], size=17,
        font=HEAD)
    text(s, 7.45, 5.65, 5.2, 0.6, "Natural. Flat. Intent-driven.", size=24, color=PINK, font=HEAD,
         align=PP_ALIGN.CENTER)
    notes(s, """
Use the AI output as a design worksheet, not a copy job:
- Main greeting -> What should the caller be able to request?
- Key mappings -> What intents do these represent?
- Call handler -> Which owner owns the outcome?
- Closed greeting -> Information, or escalation?
- Recorded FAQ -> Can this become knowledge?
- Operator transfer -> Fallback for low confidence.
Direct Inward Dial numbers from legacy Unity are ported to the AI Receptionist. Instead of WAV prompts and button mappings, admins define multilingual natural-language personas in Control Hub, with fallback and escalation driven by NLP intent detection rather than key presses.
""")

    # 20. Play 3 - AI Receptionist photo
    s = new_slide(prs, L_TWO_THIRD_PHOTO, "24/7 AI Receptionist")
    for ph in s.placeholders:
        if ph.placeholder_format.type == 18:
            ph.insert_picture(PHOTOS["receptionist"])
    eyebrow(s, "Play 3 · Webex AI Receptionist", PINK, y=0.25)
    pts = [("No-code in Control Hub", BLUE), ("Answers FAQs from your knowledge", VIOLET),
           ("Routes by intent + business hours", PINK), ("Multilingual personas", ORANGE)]
    for i, (p_, col) in enumerate(pts):
        y = 2.65 + i * 0.95
        box(s, 0.667, y + 0.08, 0.12, 0.6, fill=col, shape=MSO_SHAPE.RECTANGLE, radius=None)
        text(s, 0.95, y, 3.6, 0.8, p_, size=19, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
    notes(s, """
Webex AI Receptionist - built for corporate environments, replacing Unity Connection touch-tone auto-attendants.
- Managed through a no-code interface in Control Hub.
- Uses NLP to discern caller intent, answer FAQs from loaded knowledge bases, and route intelligently based on business hours.
- Customizable multilingual, natural-language personas and voices.
- Available 24/7 - no more "please listen carefully as our menu options have changed".
""")

    # 21. Play 4 - Dial plan
    s = new_slide(prs, L_TITLE_ONLY, "Trace the dial plan before you move it")
    eyebrow(s, "Play 4 · CUCM Dial Plan \u2192 Webex Calling", ORANGE)
    groups = [("NUMBERING", ["Directory numbers", "Shared lines"], CYAN),
              ("POLICY", ["Partitions", "Calling search spaces", "Calling privileges"], BLUE),
              ("ROUTING", ["Route patterns", "Translation patterns", "Route lists / groups"], ORANGE)]
    for i, (g, objs, col) in enumerate(groups):
        x = 0.667 + i * 4.1
        box(s, x, 1.5, 3.6, 2.85, fill=PANEL, radius=0.18)
        text(s, x + 0.25, 1.65, 3.1, 0.35, g, size=14, color=col, font=HEAD)
        for j, o in enumerate(objs):
            box(s, x + 0.25, 2.1 + j * 0.72, 3.1, 0.58, fill=WHITE, line=col, line_width=1.5, radius=0.1,
                content=o, size=17, font=HEAD)
        if i < 2:
            arrow(s, x + 3.65, 2.9, x + 4.05, 2.9, color=GREY, width=2.5)
    text(s, 0.667, 4.62, 6, 0.35, "GENAI DOES THE TRACING", size=13, color=GREY, font=HEAD)
    ai_chips = [("search_person", "Find duplicates"), ("chat", "Explain conflicts"),
                ("phone", "Normalize to E.164"), ("analytics", "Generate test plans")]
    for i, (ic, lbl) in enumerate(ai_chips):
        x = 0.667 + i * 3.05
        box(s, x, 5.0, 2.85, 1.2, fill=grad(ORANGE, PINK, 45), radius=0.16)
        icon(s, ic, x + 0.22, 5.3, 0.6, WHITE)
        text(s, x + 0.95, 5.0, 1.85, 1.2, lbl, size=18, color=WHITE, font=HEAD, anchor=MSO_ANCHOR.MIDDLE,
             spacing=0.95)
    notes(s, """
Play 4 - CUCM dial plan to Webex Calling.
A dial plan is a dependency map: numbering (DNs, devices, shared lines) feeds policy (partitions, CSS, calling privileges), which feeds routing (route patterns, translation patterns, route lists/groups, locations, voicemail profiles). A change in one object cascades.
Review sequence: DNs and devices -> shared lines and partitions -> CSS -> route and translation patterns -> route lists/groups and locations -> voicemail profiles and calling privileges.
GenAI accelerates the tracing on a sanitized BAT export: flag duplicate extensions before they cause routing failures, explain conflicts in plain language, group records for normalization/redesign/retirement, and generate E.164 normalization tables and test plans.
Control Hub migration wizards then map route patterns and DNs to Webex Calling using E.164 normalization and flat extensions.
""")

    # 22. Hybrid safety net
    s = new_slide(prs, L_TITLE_ONLY, "Hybrid routing = zero-downtime cutover")
    eyebrow(s, "Play 4 · Hybrid safeguards", ORANGE)
    nodes = {"cloud": (0.9, 1.7, "Webex Calling", "cloud", grad(SKY, BLUE, 45)),
             "lgw": (3.55, 3.2, "Local Gateway", "integration", ORANGE),
             "cucm": (6.2, 1.7, "CUCM (remaining)", "phone", DARK),
             "pstn": (3.55, 5.0, "PSTN", "remote", GREY)}
    d = 1.15
    centers = {k: (v[0] + d / 2, v[1] + d / 2) for k, v in nodes.items()}
    arrow(s, *centers["cloud"], *centers["lgw"], color=MID, width=3, head=False)
    arrow(s, *centers["cucm"], *centers["lgw"], color=MID, width=3, head=False)
    arrow(s, *centers["lgw"], *centers["pstn"], color=MID, width=3, head=False)
    arrow(s, centers["cloud"][0] + 0.6, centers["cloud"][1], centers["cucm"][0] - 0.6, centers["cucm"][1],
          color=ORANGE, width=2.25, dash="dash", tail=True)
    text(s, 2.3, 1.85, 3.8, 0.35, "Unified hybrid dial plan", size=15, color=ORANGE, font=HEAD,
         align=PP_ALIGN.CENTER)
    for k, (x, y, lbl, ic, col) in nodes.items():
        icon_disc(s, ic, x, y, d, col)
        if k == "pstn":
            text(s, x + d + 0.15, y + 0.35, 1.8, 0.45, lbl, size=17, font=HEAD)
        elif k == "lgw":
            text(s, x + d + 0.15, y + 0.35, 2.2, 0.45, lbl, size=17, font=HEAD)
        else:
            text(s, x - 0.5, y + d + 0.08, d + 1.0, 0.45, lbl, size=17, font=HEAD, align=PP_ALIGN.CENTER)
    box(s, 8.1, 1.45, 4.57, 5.0, fill=PANEL, radius=0.2)
    text(s, 8.4, 1.62, 4.0, 0.35, "CUTOVER CHECKS", size=13, color=GREY, font=HEAD)
    checks = ["Extension reachability", "PSTN in / out", "Emergency calling", "Voicemail routing",
              "Shared-line behavior", "Calling permissions", "Caller ID", "Failover / rollback"]
    for i, c in enumerate(checks):
        y = 2.05 + i * 0.53
        circle(s, 8.4, y + 0.07, 0.34, ORANGE, content="\u2713", size=13, color=WHITE, font=HEAD)
        text(s, 8.9, y, 3.7, 0.48, c, size=17, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
To prevent downtime during phased CUCM and Hunt Pilot migrations, establish Local Gateways and a unified hybrid dial plan. This lets remaining legacy extensions and newly provisioned cloud users/queues interact fluidly.
Validate before each wave: on-prem <-> cloud calls, extension reachability, PSTN ingress/egress, emergency calling, voicemail routing, shared-line behavior, calling permissions, caller ID, and a tested failover/rollback plan.
""")

    # 23. Voicemail AI
    s = new_slide(prs, L_TITLE_ONLY2, "Every voicemail, readable")
    eyebrow(s, "Play 4 · Voicemail AI Transcription", ORANGE)
    box(s, 0.667, 1.6, 6.4, 3.6, fill=WHITE, line=LINE, radius=0.2)
    icon_disc(s, "phone", 0.95, 1.85, 0.8, grad(ORANGE, PINK, 45), ratio=0.5)
    text(s, 1.95, 1.85, 4.0, 0.4, "Voicemail  ·  Dana R.", size=18, font=HEAD)
    text(s, 1.95, 2.25, 4.0, 0.35, "0:42  ·  Today 9:14 AM", size=14, color=GREY)
    icon(s, "waveform", 5.95, 1.85, 0.8, ORANGE)
    text(s, 0.95, 2.95, 5.9, 2.0, "\u201cHi, it's Dana - the router shipment lands Tuesday. Call me back to "
                                  "confirm the install window.\u201d", size=20, color=DARK,
         spacing=1.1)
    text(s, 7.6, 1.45, 1.6, 1.6, "6", size=110, font=HEAD, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.1, 1.75, 3.5, 1.1, "languages, natively", size=26, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    langs = ["EN", "FR", "DE", "IT", "JA", "ES"]
    for i, l_ in enumerate(langs):
        box(s, 7.65 + i * 0.83, 3.25, 0.72, 0.55, fill=ORANGE if i % 2 == 0 else PINK, radius=0.27, content=l_,
            size=16, color=WHITE, font=HEAD)
    text(s, 7.65, 4.1, 5.0, 1.2, ["Toggle per Location in Control Hub", "Delivered to Webex App + email"],
         size=18, font=HEAD, color=DARK, space_after=6)
    notes(s, """
Migrating the CUCM dial plan to Webex Calling unlocks native Voicemail AI Transcription.
- Enabled at the Location level in Control Hub.
- Speech-to-text delivered directly in the Webex App and via email.
- Natively supports English, French, German, Italian, Japanese and Spanish.
""")

    # 24. Shared lines
    s = new_slide(prs, L_TITLE_ONLY, "Shared lines are workflow signals")
    eyebrow(s, "Play 4 · Shared lines", ORANGE)
    text(s, 0.667, 1.35, 4.8, 0.3, "WHAT YOU SEE IN CUCM", size=13, color=GREY, font=HEAD, align=PP_ALIGN.CENTER)
    text(s, 7.2, 1.35, 5.47, 0.3, "LIKELY WEBEX PATTERN", size=13, color=GREY, font=HEAD, align=PP_ALIGN.CENTER)
    sig = [("Executive / assistant coverage", "Executive Assistant"),
           ("Informal queuing", "Call Queue or Customer Assist"),
           ("Multi-device mobility", "Multiple devices per user"),
           ("Departmental ownership", "Hunt Group or virtual line"),
           ("Legacy workaround", "Retire")]
    for i, (a, b) in enumerate(sig):
        y = 1.75 + i * 0.9
        box(s, 0.667, y, 4.8, 0.72, fill=PANEL, radius=0.12, content=a, size=18, font=HEAD, color=DARK)
        arrow(s, 5.55, y + 0.36, 7.1, y + 0.36, color=ORANGE, width=2.25)
        box(s, 7.2, y, 5.47, 0.72, fill=DARK if b == "Retire" else BLUE, radius=0.12, content=b, size=18,
            color=WHITE, font=HEAD)
    text(s, 0.667, 6.3, 12.0, 0.45, "Ask why the line is shared - then pick the pattern.", size=18, color=GREY,
         align=PP_ALIGN.CENTER)
    notes(s, """
Shared lines are rarely "just" shared lines - they're signals of a workflow:
- Executive/assistant coverage -> Webex Calling Executive Assistant.
- Informal queuing -> a real Call Queue, or Customer Assist for frontline teams.
- Multiple-device mobility -> multiple devices per user in Webex Calling.
- Departmental ownership / shared responsibility -> Hunt Group or virtual line.
- Legacy workarounds -> retire.
Have the LLM classify each shared-line appearance from the sanitized export, then confirm with the business owner. Validate the target feature fit for your tenant before you build.
""")

    # 25. Play 5 - devices
    s = new_slide(prs, L_TITLE_ONLY, "Keep the macros. Then modernize them.")
    eyebrow(s, "Play 5 · CE Endpoints \u2192 RoomOS", VIOLET)
    top = [("CE endpoint", "devices", GREY), ("Webex Device Connector", "updown", VIOLET),
           ("Control Hub Workspace", "cloud", BLUE)]
    for i, (lbl, ic, col) in enumerate(top):
        x = 0.667 + i * 4.1
        box(s, x, 1.5, 3.6, 0.9, fill=PANEL, radius=0.45)
        icon_disc(s, ic, x + 0.12, 1.57, 0.76, col, ratio=0.55)
        text(s, x + 1.0, 1.5, 2.55, 0.9, lbl, size=17, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)
        if i < 2:
            arrow(s, x + 3.65, 1.95, x + 4.05, 1.95, color=GREY, width=2.5)
    cx, cy, r = 3.05, 4.65, 1.35
    ring = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    set_fill(ring, None)
    set_line(ring, VIOLET, 3)
    _no_effects(ring)
    text(s, cx - 0.9, cy - 0.4, 1.8, 0.8, ["GenAI", "loop"], size=20, font=HEAD, color=VIOLET,
         align=PP_ALIGN.CENTER, spacing=0.9)
    loop = [("docs", "Legacy macro", 0, -1), ("ai", "LLM \u2192 xAPI", 1, 0),
            ("analytics", "Macro Editor test", 0, 1), ("download", "Deploy", -1, 0)]
    nd = 0.8
    for ic, lbl, dx, dy in loop:
        nx, ny = cx + dx * r - nd / 2, cy + dy * r - nd / 2
        icon_disc(s, ic, nx, ny, nd, grad(VIOLET, PINK, 45), ratio=0.55)
        if dx == 0:
            text(s, nx + nd + 0.1, ny + 0.15, 2.4, 0.5, lbl, size=16, font=HEAD)
        elif dx == 1:
            text(s, nx + nd + 0.1, ny + 0.15, 1.8, 0.5, lbl, size=16, font=HEAD)
        else:
            text(s, nx - 1.05, ny + 0.15, 1.0, 0.5, lbl, size=16, font=HEAD, align=PP_ALIGN.RIGHT)
    box(s, 6.9, 3.0, 5.77, 1.55, fill=WHITE, line=ORANGE, line_width=2, radius=0.18)
    text(s, 7.2, 3.15, 5.2, 0.35, "HEADS UP", size=13, color=ORANGE, font=HEAD)
    text(s, 7.2, 3.5, 5.3, 1.0, ["CE 9.x \u2194 RoomOS xAPI isn't always 1:1", "Test HVAC & AV integrations first"],
         size=18, font=HEAD, space_after=4)
    box(s, 6.9, 4.8, 5.77, 1.5, fill=grad(VIOLET, BLUE, 0), radius=0.18)
    text(s, 7.2, 4.95, 5.2, 0.35, "PRESERVED", size=13, color=WHITE, font=HEAD)
    text(s, 7.2, 5.3, 5.3, 0.9, "UI Extensions  ·  Macros  ·  Central management", size=18, color=WHITE,
         font=HEAD)
    notes(s, """
Play 5 - Collaboration Endpoints (CE) to RoomOS.
- The Webex Device Connector transfers devices and centralizes management as Workspaces in Control Hub. Existing JavaScript macros and UI Extensions are preserved.
- As devices are reorganized into unified Workspaces, legacy macros often need updates. Admins use LLMs (via Webex Connect's Evaluate node or external APIs) to translate legacy CE macro code into RoomOS-compatible scripts. Input: a natural-language description or legacy snippet. Output: modernized xAPI commands.
- Limitation: xAPI event mappings between CE 9.x and RoomOS are not always 1:1. LLM-generated JavaScript must be tested in the local Macro Editor so hardware integrations (HVAC sensors, AV controllers) still work after migration.
Assessment first: inventory and models, registration state, UI extensions and macros, xAPI integrations, workspace utilization.
""")

    # 26. Every room AI-ready (photo)
    s = new_slide(prs, L_SECTION_BODY_PHOTO)
    for ph in s.placeholders:
        f = ph.placeholder_format
        if f.type == 18:
            ph.insert_picture(PHOTOS["video_room"])
        elif f.idx == 12:
            ph.text = "PLAY 5 · ROOMOS"
        elif f.idx == 0:
            ph.text = "Every room, AI-ready"
        elif f.idx == 13:
            fill_tf(ph.text_frame, ["Webex AI Assistant", "Real-time captions", "One pane: Control Hub"],
                    size=20, font=HEAD, space_after=10)
    notes(s, """
Once devices land on RoomOS as cloud-managed Workspaces, they pick up Webex AI Assistant capabilities and real-time audio captioning, and are managed from the same Control Hub as calling and contact center.
""")

    # 27. Triage endpoints
    s = new_slide(prs, L_TITLE_ONLY, "Triage every endpoint")
    eyebrow(s, "Play 5 · Device decisions", VIOLET)
    tri = [("Supported device", "Migrate", BLUE), ("Legacy macro", "Refactor + test", VIOLET),
           ("Unsupported endpoint", "Replace / retain", ORANGE), ("Unused UI Extension", "Remove", DARK),
           ("Hard-coded integration", "Redesign", PINK), ("Duplicate naming", "Normalize", BLUE),
           ("Low utilization", "Re-evaluate", VIOLET), ("Complex control", "Validate", PINK)]
    for i, (f_, a, col) in enumerate(tri):
        x = 0.667 + (i % 4) * 3.05
        y = 1.55 + (i // 4) * 2.45
        box(s, x, y, 2.85, 2.2, fill=PANEL, radius=0.16)
        box(s, x, y, 2.85, 0.1, fill=col, shape=MSO_SHAPE.RECTANGLE, radius=None)
        text(s, x + 0.25, y + 0.35, 2.4, 0.7, f_, size=16, color=GREY, spacing=0.95)
        text(s, x + 0.25, y + 1.1, 2.45, 0.9, a, size=24, color=col, font=HEAD, spacing=0.9)
    notes(s, """
Let the LLM categorize the inventory (sanitized), identify aging hardware, summarize macros, document xAPI usage and flag hard-coded addresses. Then act:
- Supported device -> migrate. Legacy macro -> refactor and test. Unsupported endpoint -> replace or retain.
- Unused UI Extension -> remove. Hard-coded integration -> redesign. Duplicate naming -> normalize.
- Low utilization -> re-evaluate placement. Complex room control -> validate end to end.
""")

    # 28. Section - guardrails
    s = new_slide(prs, L_SECTION1, "03\u000bTrust & Guardrails")
    notes(s, "Section: how we keep AI output trustworthy.")

    # 29. Maturity staircase
    s = new_slide(prs, L_TITLE_ONLY, "Earn your way to automation")
    eyebrow(s, "AI output levels")
    lv = [("1", "Advisory", "Summaries · diagrams · questions", 1.75, CYAN),
          ("2", "Draft artifacts", "CSV · JSON · code · test plans", 2.85, BLUE),
          ("3", "Controlled automation", "APIs · schema checks · pipelines", 3.95, grad(BLUE, PINK, 45))]
    base = 6.3
    for i, (n, name, ex, hgt, col) in enumerate(lv):
        x = 0.667 + i * 4.05
        box(s, x, base - hgt, 3.85, hgt, fill=col, radius=0.18)
        text(s, x + 0.3, base - hgt + 0.1, 1.2, 1.0, n, size=54, color=WHITE, font=HEAD)
        text(s, x + 1.25, base - hgt + 0.25, 2.5, 0.8, name, size=21, color=WHITE, font=HEAD, spacing=0.9,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, base - hgt + 1.1, 3.4, 0.5, ex, size=15, color=WHITE)
    text(s, 0.667, 1.55, 7.6, 0.9, [[("Start advisory. ", {"color": INK}),
                                     ("Prove reliability", {"color": PINK}),
                                     (" before you automate.", {"color": INK})]],
         size=26, font=HEAD, spacing=0.95)
    notes(s, """
Three levels of AI output:
- Level 1, Advisory: summaries, diagrams and questions that accelerate informed decisions.
- Level 2, Draft artifacts: mapping tables, CSV/JSON, code and test plans that speed delivery.
- Level 3, Controlled automation: API retrieval, schema validation and pipeline integration for reliable execution.
Start with advisory outputs. Prove reliability before automation.
""")

    # 30. Four gates
    s = new_slide(prs, L_TITLE_ONLY, "Four gates to production")
    eyebrow(s, "Validate every AI output")
    icon_disc(s, "docs", 0.667, 3.0, 1.2, GREY)
    text(s, 0.4, 4.3, 1.75, 0.4, "AI draft", size=17, font=HEAD, align=PP_ALIGN.CENTER)
    gates = [("Syntax", "Valid format?", CYAN), ("Schema", "Right fields & types?", BLUE),
             ("Platform", "Supported in Webex?", VIOLET), ("Intent", "Right caller experience?", PINK)]
    for i, (g, q, col) in enumerate(gates):
        x = 2.35 + i * 2.25
        box(s, x, 1.6, 1.95, 4.2, fill=PANEL, radius=0.18)
        box(s, x, 1.6, 1.95, 0.12, fill=col, shape=MSO_SHAPE.RECTANGLE, radius=None)
        circle(s, x + 0.55, 2.0, 0.85, col, content="\u2713", size=30, color=WHITE, font=HEAD)
        text(s, x + 0.1, 3.1, 1.75, 0.5, g, size=22, font=HEAD, align=PP_ALIGN.CENTER)
        text(s, x + 0.15, 3.65, 1.65, 1.2, q, size=16, color=GREY, align=PP_ALIGN.CENTER, spacing=0.95)
    arrow(s, 1.9, 3.6, 2.3, 3.6, color=GREY, width=2.5)
    arrow(s, 11.3, 3.6, 11.65, 3.6, color=GREY, width=2.5)
    icon_disc(s, "cloud", 11.7, 3.0, 1.2, grad(BLUE, PINK, 45))
    text(s, 11.35, 4.3, 1.9, 0.4, "Production", size=17, font=HEAD, align=PP_ALIGN.CENTER)
    notes(s, """
Validate AI outputs through four gates:
1. Syntax - correct data format (valid JSON/CSV/JS).
2. Schema - confirm fields and types against current Webex import schemas.
3. Platform - confirm the target platform actually supports the feature/service.
4. Intent - align with business intent and the desired user experience.
Checklist examples for UCCX -> WxCC: caller paths, business hours, prompts, database/API dependencies, queue/skill behavior, RONA/agent state, timeout/error paths, simplification approval, AI Agent escalation, schema conformance.
""")

    # 31. Failure modes
    s = new_slide(prs, L_TITLE_ONLY, "Know the failure modes")
    eyebrow(s, "Risk \u2192 fix")
    fm = [("Hallucinated API", "Supply the docs"), ("Outdated terms", "Date the prompt"),
          ("Wrong mapping", "Force classification"), ("Lost leading \u201c+\u201d", "Numbers as strings"),
          ("Edited CSV headers", "Lock the headers"), ("Missing relationships", "Give a data dictionary"),
          ("PII in prompts", "Sanitize first"), ("Overconfidence", "Require labels")]
    for i, (risk, fix) in enumerate(fm):
        x = 0.667 + (i % 4) * 3.05
        y = 1.55 + (i // 4) * 2.4
        box(s, x, y, 2.85, 2.15, fill=PANEL, radius=0.16)
        text(s, x + 0.25, y + 0.25, 2.4, 0.5, risk, size=17, color=PINK, font=HEAD)
        arrow(s, x + 0.3, y + 0.95, x + 1.0, y + 0.95, color=MID, width=2)
        text(s, x + 0.25, y + 1.15, 2.45, 0.9, fix, size=22, font=HEAD, spacing=0.9)
    notes(s, """
Common failure modes and mitigations:
- Hallucinated API -> supply the documentation in context.
- Outdated terminology -> put today's date and product versions in the prompt.
- Incorrect mapping -> require explicit classification with rationale.
- Lost leading "+" on E.164 numbers -> treat numbers as strings.
- Modified CSV -> lock the headers to the Control Hub template.
- Missing relationships -> provide a data dictionary.
- PII included -> sanitize before sending.
- Overconfidence -> require confidence labels / REQUIRES_VALIDATION.
- And always: incorrect behavior is caught by running tests.
""")

    # 32. Choosing the model
    s = new_slide(prs, L_TITLE_ONLY, "Pick your model like a vendor")
    eyebrow(s, "Model selection")
    crit = [("lock", "Data protection", "Where does data go?", CYAN),
            ("docs", "Context window", "Fits the whole export?", BLUE),
            ("download", "Structured output", "Reliable JSON / CSV?", VIOLET),
            ("ai", "Reasoning & code", "Handles routing logic?", PINK),
            ("analytics", "Audit & retention", "Meets governance?", ORANGE)]
    for i, (ic, name, q, col) in enumerate(crit):
        x = 0.667 + i * 2.45
        icon_disc(s, ic, x + 0.45, 1.75, 1.4, col, ratio=0.5)
        text(s, x, 3.35, 2.3, 0.5, name, size=19, font=HEAD, align=PP_ALIGN.CENTER, wrap=False)
        text(s, x, 3.9, 2.3, 0.8, q, size=16, color=GREY, align=PP_ALIGN.CENTER)
    box(s, 2.67, 5.3, 8.0, 0.75, fill=INK, radius=0.37,
        content=[[("Rule of thumb:  ", {"color": SKY}), ("enterprise tenant or it doesn't touch config", {})]],
        size=18, color=WHITE, font=HEAD)
    notes(s, """
Selecting the right AI model:
1. Data-protection requirements - where prompts and outputs are processed and stored.
2. Context-window needs - large AEF scripts and DataDump CSVs need room.
3. Structured-output reliability - consistent, machine-readable JSON/CSV.
4. Coding and reasoning performance - routing logic and macro translation.
5. Audit and retention compliance - meet enterprise governance standards.
""")

    # 33. Evidence pack
    s = new_slide(prs, L_TITLE_ONLY, "Ship an evidence pack, not just a config")
    eyebrow(s, "Migration outputs")
    packs = [("DISCOVER", ["Inventory", "Sanitization record", "Behavior summary", "Dependency map"], CYAN),
             ("DESIGN", ["Assumptions", "Target mapping", "Exception list", "Draft artifacts"], BLUE),
             ("DELIVER", ["Validation results", "Rollback plan", "Approval", "Documentation"], PINK)]
    for i, (hdr, its, col) in enumerate(packs):
        x = 0.667 + i * 4.1
        box(s, x, 1.55, 3.8, 4.75, fill=PANEL, radius=0.18)
        box(s, x, 1.55, 3.8, 0.8, fill=col, radius=0.18, content=hdr, size=20, color=WHITE, font=HEAD)
        box(s, x, 2.05, 3.8, 0.3, fill=col, shape=MSO_SHAPE.RECTANGLE, radius=None)
        for j, it in enumerate(its):
            y = 2.7 + j * 0.88
            circle(s, x + 0.3, y + 0.08, 0.42, WHITE, line=col, content="\u2713", size=15, color=col,
                   font=HEAD)
            text(s, x + 0.9, y, 2.8, 0.6, it, size=20, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """
Every migration wave produces a package, not just a config file:
Discover - inventory, sanitization record, behavior summary, dependency map.
Design - assumptions, target mapping, exception list, draft artifacts.
Deliver - validation results, rollback plan, approval, documentation.
This is what makes AI-assisted work auditable and repeatable.
""")

    # 34. Five plays one pattern
    s = new_slide(prs, L_TITLE_ONLY, "Five plays. One pattern.")
    eyebrow(s, "Recap")
    heads = ["PLAY", "EXTRACT", "ANALYZE", "LAND ON", "AI WIN"]
    cw = [2.3, 2.3, 2.3, 2.55, 2.55]
    xs = [0.667]
    for w_ in cw[:-1]:
        xs.append(xs[-1] + w_)
    for x, w_, h_ in zip(xs, cw, heads):
        box(s, x + 0.04, 1.5, w_ - 0.08, 0.55, fill=INK, radius=0.08, content=h_, size=14, color=WHITE, font=HEAD)
    grid = [("UCCX", "AEF scripts", "Caller paths", "Contact Center", "AI Agent", BLUE),
            ("Hunt Pilots", "BAT export", "Queue behavior", "Customer Assist", "Summaries", CYAN),
            ("Unity", "DataDump CSV", "Menu tree", "Auto Attendant", "AI Receptionist", PINK),
            ("Dial Plan", "BAT export", "Conflicts", "Webex Calling", "VM transcription", ORANGE),
            ("Devices", "Inventory", "Readiness", "RoomOS", "AI Assistant", VIOLET)]
    for r, row in enumerate(grid):
        y = 2.15 + r * 0.85
        col = row[5]
        for c, (x, w_) in enumerate(zip(xs, cw)):
            if c == 0:
                box(s, x + 0.04, y, w_ - 0.08, 0.72, fill=col, radius=0.08, content=row[0], size=18, color=WHITE,
                    font=HEAD)
            elif c == 4:
                box(s, x + 0.04, y, w_ - 0.08, 0.72, fill=WHITE, line=col, line_width=2, radius=0.08,
                    content=row[4], size=17, color=col, font=HEAD)
            else:
                box(s, x + 0.04, y, w_ - 0.08, 0.72, fill=PANEL, radius=0.08, content=row[c], size=17, color=DARK)
    notes(s, """
Five plays, one pattern - extract, analyze, modernize:
- UCCX: script -> caller paths -> WxCC / AI Agent.
- Hunt Pilots: BAT -> queue behavior -> Customer Assist.
- Unity: exports -> menu hierarchy -> AI Receptionist.
- CUCM: dial plan -> conflicts -> Webex Calling.
- Devices: inventory -> readiness -> cloud-managed RoomOS.
""")

    # 35. Key takeaways
    s = new_slide(prs, L_TITLE_ONLY2, "Five things to do Monday")
    eyebrow(s, "Key takeaways")
    tk = [("Recover intent", "before mapping"), ("Sanitize data", "before AI"),
          ("Expose", "assumptions"), ("Validate", "against schemas"), ("Transform,", "don't just migrate")]
    for i, (a, b) in enumerate(tk):
        x = 0.667 + i * 2.45
        hero = i == 4
        box(s, x, 1.6, 2.25, 3.9, fill=grad(BLUE, PINK, 45) if hero else PANEL, radius=0.18)
        text(s, x + 0.25, 1.8, 1.8, 1.0, f"0{i + 1}", size=44, color=WHITE if hero else BLUE, font=HEAD)
        text(s, x + 0.25, 3.3, 1.9, 0.9, a, size=22, color=WHITE if hero else INK, font=HEAD, spacing=0.9)
        text(s, x + 0.25, 4.25, 1.9, 0.9, b, size=17, color=WHITE if hero else GREY, spacing=0.95)
    notes(s, """
1. Recover intent before mapping - understand what the legacy config was trying to do.
2. Sanitize data before AI.
3. Expose assumptions and exceptions - REQUIRES_VALIDATION is a feature.
4. Validate against current schemas - four gates, every time.
5. Transform, don't just migrate - land on AI Agent, AI Receptionist, Voicemail AI Transcription, and RoomOS AI.
""")

    # 36. Closing statement
    s = new_slide(prs, L_STATEMENT1, "Your legacy config is data.\nGenAI turns it into a roadmap.")
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 11:
            ph.text = "Extract  \u2192  Sanitize  \u2192  Analyze  \u2192  Transform  \u2192  Validate"
    remove_empty_placeholders(s)
    notes(s, "Closing message. Invite questions.")

    # 37. Resources
    s = new_slide(prs, L_TITLE_ONLY, "Keep learning")
    eyebrow(s, "Resources")
    box(s, 0.667, 1.5, 5.85, 4.85, fill=PANEL, radius=0.18)
    text(s, 0.95, 1.7, 5.3, 0.4, "CISCO LIVE SESSIONS", size=14, color=BLUE, font=HEAD)
    text(s, 0.95, 2.2, 5.4, 4.0, ["BRKCOL-2043  ·  Cisco Live 2025 US", "BRKCOL-2076  ·  Cisco Live 2026 EMEA",
                                   "BRKCOL-2530  ·  Cisco Live 2025 Melbourne", "BRKCOL-2084", "BRKCOL-2185",
                                   "BRKCCT-1512"], size=18, font=HEAD, space_after=10)
    box(s, 6.82, 1.5, 5.85, 4.85, fill=PANEL, radius=0.18)
    text(s, 7.1, 1.7, 5.3, 0.4, "DOCS & TOOLS", size=14, color=PINK, font=HEAD)
    text(s, 7.1, 2.2, 5.4, 4.0, ["Migrate Unified CM data to Webex Calling", "Webex Calling Customer Assist",
                                  "AI Receptionist for Webex Calling", "Voicemail transcription for Webex Calling",
                                  "Webex AI Agent At-a-Glance", "Webex Devices xAPI docs",
                                  "CiscoDevNet / roomdevices-macros-samples"], size=17, font=HEAD, space_after=8)
    notes(s, """
Key citations from the research:
Webex AI Agent: personalize every interaction; Webex AI Agent At-a-Glance; Cisco Collaboration Flex 3.0 Contact Center Data Sheet; UCCE 15.0(1) Features Guide - Webex AI Agent; Webex What's New Spring 2025; the AI-first front door; Webex Calling Customer Assist Solution Overview and FAQ; Understand Routing and Queueing in Webex Contact Center; Webex Calling Call Queue vs. Customer Assist feature comparison; Configure Webex Calling Customer Assist for Unified CM; Migrate from Cisco Customer Journey Platform (R10) and CC-One (R9) to Webex Contact Center; Configure Call Queuing (CUCM 12.5); AI Receptionist for Webex Calling (and At-a-Glance); Collaboration Flex Plan 3.0 and Webex Suite Data Sheet; Voicemail transcription for Webex Calling; Troubleshoot Voicemail Issues in Webex Calling; Configure and manage voicemail settings for a Webex Calling user; Configure Unified CM for Webex Calling; Migrate Unified CM data to Webex Calling; Webex Calling Migration Journeys; Transitioning from Unified CM to Webex Calling; Manage auto attendants; Cisco Unity Connection 15 System Administration Guide - Tools; CUPI API - Call Handler; Webex Attendant Console overview and FAQ; RoomOS for Cisco devices; CE 9.9 / 9.14 Customization Guides; Cisco Desk, Board and Room Series WLAN Deployment Guide; CiscoDevNet/roomdevices-macros-samples; Webex Devices xAPI docs.
Cisco Live: BRKCOL-2043 (2025 US), BRKCOL-2076 (2026 EMEA), BRKCOL-2084, BRKCOL-2185, BRKCOL-2530 (2025 Melbourne), BRKCCT-1512, CLS-2221 (Customer Assist).
Internal partner enablement (not for public slides): PIW and EMEA Collab PVT sessions on Customer Assist, One-Click UCCX Migration, Webex Suite Cloud & AI, Calling Hybrid & Cloud-Connected UC, Calling Migration Bootcamp / Setup Assist, Programmability Workshops.
""")

    # 38. Thank you
    s = new_slide(prs, L_THANKS, "Thank you")
    notes(s, "Thank the audience. Q&A.")

    prs.save(OUT)
    print("saved", OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    build()
