"""Build the WebexOne 2026 technical session deck on the official template."""
import os
import re

import cairosvg
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x0A, 0x0A, 0x0A)
DARK = RGBColor(0x28, 0x28, 0x28)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
MID = RGBColor(0xB4, 0xB9, 0xC0)
LINE = RGBColor(0xD6, 0xD6, 0xD6)
PANEL = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CYAN = RGBColor(0x06, 0xA8, 0xFF)
BLUE = RGBColor(0x0A, 0x60, 0xFF)
VIOLET = RGBColor(0x8A, 0x3F, 0xFC)
PINK = RGBColor(0xFF, 0x00, 0x7F)
ORANGE = RGBColor(0xFF, 0x90, 0x00)
SKY = RGBColor(0x00, 0xC8, 0xFF)

HEAD = "Inter Semi Bold"
BODY = "Inter"
MONO = "Consolas"

# ----------------------------------------------------------------- helpers
def hexs(c):
    return str(c)


ICON_ALPHA = "/home/ubuntu/work2/icons_alpha"
ICON_CACHE = "/home/ubuntu/work2/icons"
ICONS = {"search_person": 0, "lock": 1, "analytics": 2, "link": 3, "waveform": 4, "devices": 5, "docs": 6,
         "mic": 7, "transcript": 8, "apps": 9, "download": 10, "cc_ai": 11, "phone": 12, "graph": 13,
         "audio": 14, "ai": 15, "chat": 16, "automation": 17, "integration": 18, "updown": 19,
         "privacy": 20, "translate": 21, "secure_cycle": 22, "remote": 23, "headset": 24, "people": 25}


def icon_png(name, color):
    from PIL import Image
    os.makedirs(ICON_CACHE, exist_ok=True)
    out = f"{ICON_CACHE}/{name}_{hexs(color)}.png"
    if not os.path.exists(out):
        alpha = Image.open(f"{ICON_ALPHA}/{ICONS[name]:02d}.png").convert("L")
        im = Image.new("RGBA", alpha.size, "#" + hexs(color))
        im.putalpha(alpha)
        im.save(out)
    return out


def icon(slide, name, x, y, size, color=INK):
    return slide.shapes.add_picture(icon_png(name, color), Inches(x), Inches(y), Inches(size), Inches(size))


def _no_effects(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        eff = style.find(qn("a:effectRef"))
        if eff is not None:
            eff.set("idx", "0")


def set_fill(shape, fill):
    if fill is None:
        shape.fill.background()
    elif isinstance(fill, tuple) and fill[0] == "grad":
        _, c1, c2, angle = fill
        shape.fill.gradient()
        shape.fill.gradient_angle = angle
        stops = shape.fill.gradient_stops
        stops[0].position, stops[1].position = 0.0, 1.0
        stops[0].color.rgb, stops[1].color.rgb = c1, c2
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill


def set_line(shape, line, width=1.0, dash=None):
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
        if dash:
            ln = shape.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", dash)


def style_runs(p, runs, size, color, font, align, spacing=None):
    if align is not None:
        p.alignment = align
    if spacing:
        p.line_spacing = spacing
    for r in runs:
        if isinstance(r, str):
            r = (r, {})
        text, o = r
        run = p.add_run()
        run.text = text
        f = run.font
        f.size = Pt(o.get("size", size))
        f.name = o.get("font", font)
        f.bold = False
        f.italic = o.get("italic", False)
        c = o.get("color", color)
        if c is not None:
            f.color.rgb = c
        if o.get("strike"):
            run._r.get_or_add_rPr().set("strike", "sngStrike")


def fill_tf(tf, content, size=18, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            spacing=None, space_after=0, margins=(0, 0, 0, 0), wrap=True):
    tf.clear()
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        if child.tag in (qn("a:spAutoFit"), qn("a:normAutofit"), qn("a:noAutofit")):
            bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    if isinstance(content, (str, tuple)):
        content = [content]
    for i, para in enumerate(content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(para, dict):
            runs = para["runs"]
            style_runs(p, runs if isinstance(runs, list) else [runs], para.get("size", size),
                       para.get("color", color), para.get("font", font), para.get("align", align),
                       para.get("spacing", spacing))
            p.space_before = Pt(para.get("before", 0))
            p.space_after = Pt(para.get("after", space_after))
        else:
            style_runs(p, para if isinstance(para, list) else [para], size, color, font, align, spacing)
            p.space_after = Pt(space_after)
    return tf


def text(slide, x, y, w, h, content, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, content, **kw)
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12,
        content=None, line_width=1.0, dash=None, **kw):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(s, fill)
    set_line(s, line, line_width, dash)
    _no_effects(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    kw.setdefault("margins", (0.12, 0.06, 0.12, 0.06))
    kw.setdefault("anchor", MSO_ANCHOR.MIDDLE)
    kw.setdefault("align", PP_ALIGN.CENTER)
    fill_tf(s.text_frame, content if content is not None else "", **kw)
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=2.0, head=True, dash=None, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)
    if tail:
        he = etree.SubElement(ln, qn("a:headEnd"))
        he.set("type", "triangle"), he.set("w", "med"), he.set("len", "med")
    if head:
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle"), te.set("w", "med"), te.set("len", "med")
    return c


def circle(slide, x, y, d, fill, line=None, content=None, **kw):
    return box(slide, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL, radius=None, content=content,
               margins=(0, 0, 0, 0), **kw)


def icon_disc(slide, name, x, y, d, fill, icon_color=WHITE, ratio=0.56, line=None):
    circle(slide, x, y, d, fill, line=line)
    s = d * ratio
    icon(slide, name, x + (d - s) / 2, y + (d - s) / 2, s, icon_color)


def eyebrow(slide, label, color=BLUE, y=0.28):
    text(slide, 0.667, y, 11.5, 0.34, label.upper(), size=14, color=color, font=HEAD)


def set_title(slide, title):
    slide.shapes.title.text = title


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def grad(c1, c2, angle=0):
    return ("grad", c1, c2, angle)


def remove_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def new_slide(prs, layout_idx, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    if title is not None and s.shapes.title is not None:
        set_title(s, title)
    return s


def pipeline(slide, nodes, y, h=2.05, x0=0.667, total=12.0, gap=0.42, label_size=18, sub_size=14):
    n = len(nodes)
    w = (total - gap * (n - 1)) / n
    for i, nd in enumerate(nodes):
        x = x0 + i * (w + gap)
        hero = nd.get("hero")
        fill = nd.get("fill", PANEL)
        fg = WHITE if hero else INK
        sub_c = WHITE if hero else GREY
        box(slide, x, y, w, h, fill=fill, line=None if hero else LINE, radius=0.14)
        icon(slide, nd["icon"], x + w / 2 - 0.33, y + 0.2, 0.66, WHITE if hero else nd.get("icon_color", BLUE))
        text(slide, x + 0.1, y + 0.95, w - 0.2, 0.5, nd["label"], size=label_size, color=fg, font=HEAD,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)
        if nd.get("sub"):
            text(slide, x + 0.1, y + 1.45, w - 0.2, h - 1.5, nd["sub"], size=sub_size, color=sub_c,
                 align=PP_ALIGN.CENTER, spacing=0.95)
        if i < n - 1:
            arrow(slide, x + w + 0.06, y + h / 2, x + w + gap - 0.06, y + h / 2, color=GREY, width=2.25)
    return w


